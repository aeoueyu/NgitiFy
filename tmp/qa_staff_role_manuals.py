from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

root = Path("output/documents/staff_role_manuals")
expected = {
    "NGITIFY_System_Administrator_User_Manual.docx": ["System Configuration", "Database Backups", "Integrity Checks", "System Audit Trail"],
    "NGITIFY_Owner_Dentist_User_Manual.docx": ["My Schedule", "My Patients", "Interactive Odontogram", "Material Usage"],
    "NGITIFY_Branch_Manager_User_Manual.docx": ["Branch Manager Dashboard", "Branch Analytics", "Branch Inventory"],
    "NGITIFY_Dentist_User_Manual.docx": ["Dentist Dashboard", "Treatment Logs", "Radiograph", "Material Usage"],
    "NGITIFY_Secretary_User_Manual.docx": ["Front Desk Dashboard", "Check In", "Add a New Patient", "Read-Only Mode"],
}

for filename, required in expected.items():
    path = root / filename
    assert path.exists() and path.stat().st_size > 20000, path
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "SCREENSHOT TO ADD" in text
    assert "NgitiFy" in text
    assert "\ufffd" not in text
    assert not any(bad in text for bad in ("\u00e2\u20ac", "\u00c3", "\u00ef\u00bf"))
    for phrase in required:
        assert phrase.lower() in text.lower(), (filename, phrase)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    screenshot_count = text.count("SCREENSHOT TO ADD")
    step_num_ids = []
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "List Number":
            continue
        num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
        assert num_pr is not None and num_pr.numId is not None
        step_num_ids.append(num_pr.numId.val)
    assert len(set(step_num_ids)) == screenshot_count, (filename, len(set(step_num_ids)), screenshot_count)
    print(f"{filename}|paragraphs={len(doc.paragraphs)}|tables={len(doc.tables)}|headings={len(headings)}|screenshots={screenshot_count}|numbered_lists={len(set(step_num_ids))}")
