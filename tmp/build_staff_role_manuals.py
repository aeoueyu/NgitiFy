from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("output/documents/staff_role_manuals")
OUT_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "123B5D"
BLUE = "2E74B5"
CYAN = "25A9D6"
PALE_BLUE = "E8F3FA"
PALE_GRAY = "F4F6F8"
MID_GRAY = "667784"
DARK = "243746"
WHITE = "FFFFFF"
AMBER = "A86400"
PALE_AMBER = "FFF4D6"
RED = "A33A3A"
PALE_RED = "FDECEC"
GREEN = "237A57"

ROLE_FILES = {
    "System Administrator": "NGITIFY_System_Administrator_User_Manual.docx",
    "Owner-Dentist": "NGITIFY_Owner_Dentist_User_Manual.docx",
    "Branch Manager": "NGITIFY_Branch_Manager_User_Manual.docx",
    "Dentist": "NGITIFY_Dentist_User_Manual.docx",
    "Secretary": "NGITIFY_Secretary_User_Manual.docx",
}


def proc(title, goal, steps, screenshot, result, note=None, warning=False):
    return {
        "title": title,
        "goal": goal,
        "steps": steps,
        "screenshot": screenshot,
        "result": result,
        "note": note,
        "warning": warning,
    }


COMMON = [
    proc(
        "Activate a Newly Created Account",
        "Set the first password using the activation message sent to the registered email address.",
        [
            "Open the NgitiFy account activation email.",
            "Select the activation link. If it has expired, ask an authorized administrator or manager to resend the invitation.",
            "Enter a password that satisfies the password rules shown on screen, then enter it again for confirmation.",
            "Submit the form and wait for the activation success message.",
            "Return to the NgitiFy login page and sign in with the registered email address and new password.",
        ],
        "Activation email, password setup form, and success message",
        "The account is activated and the user can proceed to the correct role dashboard.",
        "Activation links and passwords must not be shared.",
    ),
    proc(
        "Sign In",
        "Access the authorized NgitiFy workspace.",
        [
            "On a desktop or laptop, open the clinic's NgitiFy web address in a current browser.",
            "Select Login.",
            "Enter the registered email address and password.",
            "Select Sign In.",
            "Verify that the role-specific dashboard and sidebar appear.",
        ],
        "Login page and role dashboard",
        "NgitiFy opens the dashboard assigned to the authenticated role.",
        "If a menu is missing or read-only, the role may not have permission for that function.",
    ),
    proc(
        "Reset a Forgotten Password",
        "Recover account access using email verification.",
        [
            "On the login page, select Forgot Password.",
            "Enter the registered email address and request a verification code.",
            "Open the email message and copy the one-time verification code.",
            "Enter the valid code in NgitiFy.",
            "Enter and confirm the new password, then submit the form.",
            "Sign in using the new password.",
        ],
        "Forgot Password, OTP verification, and New Password screens",
        "The password is changed and the login page becomes available.",
        "Use only the most recent code and do not disclose it to another person.",
    ),
    proc(
        "Use the Sidebar and Dashboard",
        "Move between authorized modules and review at-a-glance information.",
        [
            "Select the sidebar arrow to expand the navigation labels.",
            "Select Dashboard to return to the role overview.",
            "Review statistic cards, alerts, appointment summaries, activity panels, and calendar information available to the role.",
            "Select a clickable card or sidebar item to open its detailed module.",
            "Use the browser Back button only when needed; prefer the sidebar to keep the current workspace context clear.",
        ],
        "Expanded sidebar and dashboard overview",
        "The selected module opens without changing the user's authorized role scope.",
    ),
    proc(
        "Manage the Notification Inbox",
        "Review staff notifications and maintain read status.",
        [
            "Select Notifications from the sidebar.",
            "Use the search box, type filter, read-status filter, or date range to narrow the list.",
            "Select a notification to open its complete message.",
            "Mark the message as read or unread as needed.",
            "Select Mark All as Read when all visible notifications have been reviewed.",
        ],
        "Notifications list, filters, and notification detail modal",
        "Notification read indicators and the unread badge update.",
    ),
    proc(
        "Use the Staff AI Assistant",
        "Ask operational questions using only the data available to the signed-in role.",
        [
            "Select the floating NgitiBot button.",
            "Choose a suggested prompt or enter a concise question about authorized clinic operations.",
            "Review the response and verify important facts in the corresponding NgitiFy module.",
            "Close the assistant when finished.",
        ],
        "Staff AI Assistant panel with a sample authorized prompt",
        "The assistant displays guidance based on the user's role and authorized data scope.",
        "AI output is decision support only. Verify clinical, financial, inventory, and account actions before acting.",
        True,
    ),
    proc(
        "Review Activity Logs",
        "Find and inspect account activity recorded for the user or authorized scope.",
        [
            "Select Activity Logs from the sidebar.",
            "Search by action, detail, or category, or apply a date range.",
            "Select View Details for the desired record.",
            "Review the recorded date, time, category, action, and complete details.",
            "Use Export CSV when an authorized offline copy is required.",
        ],
        "Activity Logs filters, table, and details modal",
        "Matching activity records and their full details are displayed.",
        "Exported files may contain sensitive operational information; store them securely.",
    ),
    proc(
        "Update Profile, Email, Password, and Preferences",
        "Maintain the signed-in user's account information and security settings.",
        [
            "Select My Profile to review personal and professional information.",
            "Select Edit Profile, update permitted fields, and save after checking the information.",
            "To change the email address, select Change Email, enter the new email and current password, then complete the verification sent by NgitiFy.",
            "Select Settings, open Account Security, enter the current password and the new password twice, then submit.",
            "In Notification Settings, turn available preferences on or off and save.",
        ],
        "My Profile edit mode, Change Email modal, and Settings page",
        "Saved profile information and account preferences persist after the page is reopened.",
    ),
    proc(
        "Log Out Securely",
        "End the current NgitiFy session.",
        [
            "Finish or save any open form.",
            "Select Logout at the bottom of the sidebar.",
            "Confirm the logout request.",
            "Verify that the login page appears.",
        ],
        "Logout confirmation and returned login page",
        "The authenticated session ends.",
        "Always log out on shared clinic computers.",
    ),
]


SCHEDULE = [
    proc(
        "Search, Filter, and Review the Schedule",
        "Locate appointments, phone calls, or walk-ins in the authorized schedule.",
        [
            "Select Schedule from the sidebar.",
            "Search by patient, dentist, branch, or procedure.",
            "Filter by source type, status, and date range as available.",
            "Select View on a schedule entry to open Schedule Details.",
            "Review the patient, dentist, branch, procedure, date, time, status, source, and operational notes.",
        ],
        "Schedule search/filter controls, results, and Schedule Details panel",
        "Only schedule entries matching the selected criteria are displayed.",
    ),
    proc(
        "Add a Schedule Entry",
        "Create an appointment, phone-call booking, or walk-in entry when authorized.",
        [
            "Open Schedule and select Add Schedule Entry.",
            "Choose the source type available to the role: Appointment, Phone Call, or Walk-in.",
            "Select the branch when a branch selector is shown.",
            "Search for and select the patient, or enter guest information when the selected source permits it.",
            "Select the dentist, date, procedure, and available time; enter contact information and notes when required.",
            "Review the information and submit the form.",
        ],
        "Add Schedule Entry form with the relevant source type",
        "NgitiFy confirms that the appointment or schedule entry was created and displays it in the schedule.",
        "Never create duplicate appointments when an existing entry only needs to be rescheduled.",
    ),
    proc(
        "Edit or Reschedule an Entry",
        "Correct schedule details without creating a duplicate record.",
        [
            "Find the schedule entry and select Update Schedule Entry.",
            "Choose the available update mode, such as full edit or reschedule.",
            "Change only the necessary fields, including dentist, procedure, date, time, contact information, or notes.",
            "Review availability and the displayed status.",
            "Save and confirm the update.",
        ],
        "Update Schedule Entry form and confirmation",
        "The existing schedule entry is updated successfully.",
    ),
    proc(
        "Update Appointment Status and Complete a Visit",
        "Move an appointment through the clinic workflow and capture completed treatment details.",
        [
            "Find the appointment in Schedule.",
            "Use the available action to Confirm, Cancel, Mark as In Clinic, or Mark as Complete.",
            "Confirm the status change when prompted.",
            "When completing a visit, enter the procedure performed, treatment category, tooth number or area, amount charged, amount paid, and next appointment when applicable.",
            "Submit and verify the updated status and treatment information.",
        ],
        "Schedule status actions and Mark Schedule as Complete form",
        "The appointment status changes and completed-visit information is recorded.",
        "Status changes affect queue, dashboard, notification, and patient records. Confirm the correct patient first.",
        True,
    ),
]


PATIENT_SEARCH = proc(
    "Search and View Patients",
    "Locate an authorized patient record and open its available details.",
    [
        "Select Manage Patients from the sidebar.",
        "Search by patient name, email, identifier, or other available criteria.",
        "Apply branch, status, or sorting controls when available.",
        "Select View Patient EMR for the correct patient.",
        "Review only the information required for the current task.",
    ],
    "Manage Patients search results and View Patient EMR action",
    "The selected patient's authorized record is displayed.",
    "Patient information is confidential and must not be opened without a work-related need.",
)

PATIENT_ADMIN = [
    proc(
        "Add a New Patient",
        "Create a patient account and clinic record when the role has edit permission.",
        [
            "Open Manage Patients and select Add New Patient.",
            "Enter all required identity, contact, birth, gender, and address information.",
            "Assign the correct branch and complete any available consent or patient-intake fields.",
            "Check for duplicate warnings before submitting.",
            "Submit the form and review the success confirmation.",
        ],
        "Add New Patient form, duplicate warning, and success message",
        "The patient is added and appears in Manage Patients; the account follows the configured activation process.",
        "Search first to prevent duplicate patient records.",
        True,
    ),
    proc(
        "Edit Patient Quick Details",
        "Correct permitted patient profile information.",
        [
            "Find the patient in Manage Patients.",
            "Select Edit Quick Details.",
            "Update only verified information in the available fields.",
            "Review required fields and duplicate warnings.",
            "Save and verify the update success message.",
        ],
        "Edit Patient form and update confirmation",
        "The patient's profile is updated successfully.",
    ),
]

PATIENT_LIFECYCLE = [
    proc(
        "Manage Patient Account Status",
        "Activate, deactivate, archive, or restore a patient when authorized.",
        [
            "Find the patient and review the current account status.",
            "Select the appropriate lifecycle action.",
            "Read the confirmation message and any dependency or retention warning.",
            "Enter a reason when required and confirm the action.",
            "Verify the updated status in the patient list or lifecycle history.",
        ],
        "Patient lifecycle actions, reason field, and updated status",
        "The patient's lifecycle status changes and an audit record is preserved.",
        "Deactivate or archive instead of deleting when records must be retained. Permanent deletion belongs to the administrator review process.",
        True,
    ),
    proc(
        "Transfer a Patient to Another Branch",
        "Change the patient's branch assignment while preserving history.",
        [
            "Find the active patient in Manage Patients.",
            "Select Transfer Branch.",
            "Choose the target branch.",
            "Enter a clear transfer reason and review the affected assignment.",
            "Confirm the transfer and verify the new branch shown in the patient record.",
        ],
        "Patient Branch Transfer form and updated branch assignment",
        "The patient is assigned to the selected target branch and the transfer is logged.",
        "Coordinate ongoing appointments and records with both branches before transfer.",
        True,
    ),
]


STAFF_MANAGEMENT = [
    proc(
        "Search and Filter Staff",
        "Locate dentists, secretaries, or other staff types available to the role.",
        [
            "Select Manage Staffs from the sidebar.",
            "Choose the relevant staff tab.",
            "Search by name or email.",
            "Filter by Active, Needs Activation, Inactive, Archived, All, or branch where available.",
            "Select View Profile for the correct record.",
        ],
        "Manage Staffs tabs, search, status filters, and results",
        "Only matching staff records are displayed.",
    ),
    proc(
        "Add a New Staff Account",
        "Register an authorized staff member and send the activation invitation.",
        [
            "Open Manage Staffs and select the correct role tab.",
            "Select Add New for that role.",
            "Enter all required personal, contact, professional, and branch-assignment information.",
            "Review the email address and assigned branch carefully.",
            "Submit the form and verify the role-specific success confirmation and activation-email notice.",
        ],
        "Role-specific Add New Staff form and activation success confirmation",
        "The staff account is created in Needs Activation or the configured initial status.",
        "Create only the minimum role needed for the staff member's job.",
        True,
    ),
    proc(
        "Edit Staff Information",
        "Correct staff details or branch assignments within the user's authority.",
        [
            "Find the staff member in the correct Manage Staffs tab.",
            "Select Edit Profile.",
            "Update the verified fields and branch assignment permitted for the role.",
            "Review professional identifiers and contact details.",
            "Save and verify the update confirmation.",
        ],
        "Edit Staff form and successful update message",
        "The staff profile displays the saved information.",
    ),
    proc(
        "Manage Staff Account Lifecycle",
        "Activate, deactivate, archive, or restore staff accounts when authorized.",
        [
            "Find the staff member and confirm the current status.",
            "Select Activate, Deactivate, Archive, or Restore as appropriate.",
            "Read the confirmation and dependency information.",
            "Enter a reason when required and confirm.",
            "Verify the new status and lifecycle history.",
        ],
        "Staff lifecycle controls, confirmation, and updated status",
        "The staff account status changes and the action is audited.",
        "Deactivate access promptly when a staff member should no longer sign in. Do not archive an account needed for active schedules or records without review.",
        True,
    ),
]


INVENTORY_VIEW = proc(
    "Search and Review Inventory",
    "Check stock by item, category, brand, batch, or branch.",
    [
        "Select Inventory from the sidebar when the module is available.",
        "Use search and the available category, brand, branch, status, or batch filters.",
        "Review available quantity, reorder level, batch, expiration, and low-stock indicators.",
        "Open an item or batch only when more detail is required.",
    ],
    "Inventory Tracker filters, stock table, and low-stock indicator",
    "Only matching inventory records are displayed.",
)

INVENTORY_EDIT = [
    proc(
        "Add or Edit an Inventory Item",
        "Maintain the inventory catalog and reorder settings.",
        [
            "Open Inventory.",
            "To create an item, select Add New Item; to modify one, select Edit for the correct item.",
            "Enter or update required item name, category, brand, unit, reorder level, branch, and other displayed fields.",
            "Review the information and save.",
            "Verify that the new or updated item appears in the tracker.",
        ],
        "Add/Edit Inventory Item form and saved inventory row",
        "The inventory item is created or updated successfully.",
    ),
    proc(
        "Receive Stock",
        "Add a stock batch and update the available quantity.",
        [
            "Find the inventory item and select Add Stock or Receive Stock.",
            "Enter the received quantity, batch or lot information, expiration date, supplier information, and other required fields.",
            "Assign the correct branch when applicable.",
            "Review the batch information and submit.",
            "Verify the new batch and updated available quantity.",
        ],
        "Receive Stock form, new batch, and updated quantity",
        "The stock batch is recorded and the item quantity increases.",
    ),
    proc(
        "Delete a Stock Batch",
        "Remove an incorrect batch record when authorized.",
        [
            "Open the item's batch details.",
            "Confirm the exact batch number, branch, quantity, and expiration date.",
            "Select Delete Batch.",
            "Review the warning and confirm only if deletion is appropriate.",
            "Verify that the batch is removed and the quantity is recalculated.",
        ],
        "Batch details and Delete Batch confirmation",
        "The selected stock batch is removed and inventory totals update.",
        "This changes stock records. Use only for an erroneous batch; preserve legitimate receiving and usage history.",
        True,
    ),
]


BRANCHES = [
    proc(
        "Add or Edit a Branch",
        "Maintain clinic branch information.",
        [
            "Select Branches from the sidebar.",
            "To create a branch, select Add Branch; to modify one, select Edit on the correct branch.",
            "Enter or update required branch name, address, contact information, operating details, and manager assignment when shown.",
            "Review the information and save.",
            "Verify that the branch appears with the correct details.",
        ],
        "Branches list and Add/Edit Branch form",
        "The branch is created or updated successfully.",
    ),
    proc(
        "Change Branch Status",
        "Activate or deactivate a branch safely.",
        [
            "Open Branches and locate the correct branch.",
            "Review active appointments, assigned staff, patients, and inventory dependencies.",
            "Select the available status action.",
            "Read the warning, provide a reason if requested, and confirm.",
            "Verify the updated branch status.",
        ],
        "Branch status action, dependency warning, and updated status",
        "The selected branch status is updated.",
        "Do not deactivate a branch until affected schedules, staff, patients, and stock are handled.",
        True,
    ),
    proc(
        "Review Branch Analytics",
        "Analyze appointments, trends, and procedure distribution for an authorized branch.",
        [
            "Open Branches and select View or Analytics for the desired branch, or open Branch Analytics when provided directly.",
            "Set From and To dates when a date range is needed.",
            "Review appointment totals, monthly trend, and procedure distribution.",
            "Clear filters to return to the default period.",
            "Export CSV or PDF when an authorized report is required.",
        ],
        "Branch Analytics date range, summary cards, and charts",
        "Analytics for the selected authorized branch and period are displayed.",
        "Analytics are operational summaries and should be checked against source records before formal reporting.",
    ),
]


CLINICAL = [
    proc(
        "Open an Assigned Patient's EMR",
        "Access clinical records from the dentist-scoped patient list.",
        [
            "Select Manage Patients or My Patients, depending on the role workspace.",
            "Search for the correct assigned patient.",
            "Select View Patient EMR.",
            "Confirm the patient's name and identifying information before editing.",
            "Use the Patient Profile, Medical and Dental History, Treatment History, Odontogram, and Radiograph Images tabs.",
        ],
        "Assigned patient list and EMR tab navigation",
        "The selected patient's authorized clinical record opens.",
        "For Owner-Dentist accounts, use My Patients for dentist editing tools; the general Manage Patients view is for clinic-wide management and may be read-only clinically.",
    ),
    proc(
        "Update Medical and Dental History",
        "Record verified health information needed for dental care.",
        [
            "Open the patient's EMR and select Medical and Dental History.",
            "Select Edit.",
            "Update the reason for consultation, previous dental care, treatment reactions, physician information, medical treatment, surgery or hospitalization, medication, allergies, bleeding time, blood type, blood pressure, conditions, and notes as applicable.",
            "Complete any conditional detail field that appears after a Yes answer.",
            "Review the information with the patient and save.",
        ],
        "Medical and Dental History edit form and saved record",
        "The verified medical and dental history is saved in the patient's EMR.",
        "Do not infer medical information. Record what was verified and escalate urgent concerns according to clinic policy.",
        True,
    ),
    proc(
        "Add and Maintain Treatment Logs",
        "Record a completed procedure, charges, payments, and follow-up details.",
        [
            "Open Treatment History in the patient's EMR.",
            "Select Add Treatment Log.",
            "Enter the procedure date, procedure name, category, branch, tooth number or area, amount charged, amount paid, next appointment, and clinical notes as applicable.",
            "Review the calculated balance and submit.",
            "Use the available edit, notes, archive, or delete action only for a documented correction and provide a reason when required.",
        ],
        "Add Treatment Log form and saved Treatment History row",
        "The treatment log appears in the patient's EMR with the recorded financial and follow-up details.",
        "Clinical and financial records must match the actual completed visit.",
        True,
    ),
    proc(
        "Update the Interactive Odontogram",
        "Record the patient's current tooth condition, surface, or treatment stage.",
        [
            "Open the Odontogram tab.",
            "Select the correct tooth using the FDI numbering shown by the system.",
            "Choose the supported condition, affected surface, treatment, or stage.",
            "Add the clinical note or workflow information requested on screen.",
            "Review the selected tooth and save.",
            "Check Odontogram History to confirm the update and its author.",
        ],
        "Interactive Odontogram tooth selection, editor, and history",
        "The patient's current odontogram and history show the saved tooth data.",
        "Confirm tooth number and surface before saving; odontogram changes become part of the clinical record.",
        True,
    ),
    proc(
        "Upload and Review a Radiograph",
        "Attach a radiograph and complete the dentist-led review workflow.",
        [
            "Open Radiograph Images and select Upload Radiograph.",
            "Choose the radiograph type, date taken, optional radiograph number, and image file; keep the image within the displayed upload limit (currently 3 MB).",
            "Enter findings or impression and optional notes, then upload.",
            "Open the image in Radiograph Review and use zoom, fit, brightness, contrast, comparison, or Auto Improve as needed.",
            "If analysis is enabled, run Analyze Radiograph and independently review its output.",
            "Record dentist findings, choose the FDI tooth or marked area, optionally link to the odontogram or treatment, and save.",
            "Generate or revise the summary, manually review it, and approve only after confirming accuracy.",
        ],
        "Radiograph upload, image tools, dentist finding, and approved summary",
        "The radiograph, dentist findings, and approved patient explanation are stored in the EMR.",
        "AI analysis and image enhancement do not replace clinical interpretation. The dentist remains responsible for the final finding and summary.",
        True,
    ),
    proc(
        "Log Material Usage",
        "Record materials consumed during a completed procedure.",
        [
            "Select Material Usage from the sidebar.",
            "Select Log New Entry.",
            "Choose the related completed appointment or patient and procedure.",
            "Add each inventory item used and its quantity; add another row for additional materials.",
            "Review branch, patient, procedure, item, and quantity, then save.",
            "Verify the new usage log and the corresponding inventory reduction.",
        ],
        "Log New Material Usage form and saved usage details",
        "The usage log is saved and the related inventory quantities update.",
        "Enter actual quantities only. Incorrect entries affect stock availability and audit records.",
        True,
    ),
]


ADMIN_ONLY = [
    proc(
        "Configure Roles and Permissions",
        "Set module access for configurable staff roles and manage delegated administrator access.",
        [
            "Open Roles and Permissions from the appropriate administration entry point.",
            "Review the matrix for Branch Manager, Dentist, and Secretary.",
            "For each module, choose the permitted level shown by the system, such as full access, read-only, or no access.",
            "Save the permission changes and have the affected user sign in again if required.",
            "Use Grant Admin Access only for a verified person who requires system-administration authority.",
        ],
        "Role Permission Matrix and Grant Admin Access controls",
        "The configured role permissions are saved and enforced in the affected workspace.",
        "Apply least privilege. Administrator access exposes system-wide configuration, audit, backup, and integrity tools.",
        True,
    ),
    proc(
        "Update System Configuration",
        "Maintain clinic information, appointment rules, templates, features, and session controls.",
        [
            "Select System Config from the sidebar.",
            "Review Clinic Information and update the clinic name, contact number, email, and address as required.",
            "Review Appointment Settings, including maximum appointments per day, allowed time slots, online booking procedures, and clinic procedures.",
            "Review account-activation and appointment-reminder email templates.",
            "Set authorized feature toggles and the session-timeout duration.",
            "Review all changes and select Save Configuration.",
        ],
        "System Configuration sections and Save Configuration action",
        "The system configuration is saved and used by the affected NgitiFy functions.",
        "Configuration changes can affect every branch and user. Record the reason and coordinate deployment-sensitive changes.",
        True,
    ),
    proc(
        "Update Website Content and Media",
        "Maintain the public website text, links, services, locations, and images.",
        [
            "Open System Config and navigate to Website Content.",
            "Update branding and social links, Home content, About and Locations content, Services, Contact, or Appointment-page content as needed.",
            "For media, paste an approved image URL or upload an approved file; use Reset only when the default should be restored.",
            "Add, edit, reorder, or remove service highlights carefully.",
            "Use Preview and check desktop presentation before saving.",
            "Save the configuration and verify the public website.",
        ],
        "Website Content editor, media controls, and website preview",
        "The approved website content is saved and appears on the public site.",
        "Use only approved copy and media. Check privacy, copyright, contact details, and booking information before publication.",
        True,
    ),
    proc(
        "Review the System Audit Trail",
        "Investigate system-wide actions and recorded details.",
        [
            "Select Audit Trail from the sidebar.",
            "Search by action, user, or detail, and apply category, role, or date filters.",
            "Select View Details for the target event.",
            "Review Date, Time, User, Role, Action, and Recorded Details.",
            "Export CSV only for an authorized audit or investigation.",
        ],
        "System Audit Logs filters, results, and Audit Log Details modal",
        "The selected audit event and its complete recorded details are displayed.",
        "Audit records are sensitive and should not be altered or distributed outside authorized review.",
    ),
    proc(
        "Review Archived Accounts",
        "Search archived records and decide whether to restore or review permanent deletion eligibility.",
        [
            "Select Archive Review from the sidebar.",
            "Search by name, email, role, or archive reason.",
            "Filter by record type and review state: Ready to Delete, Waiting Retention, or Blocked by History.",
            "To recover an account, select Restore and confirm.",
            "For permanent deletion review, open the blocker details and confirm that retention and linked-history conditions are satisfied before proceeding.",
        ],
        "Archive Review filters, Restore action, and permanent-delete blocker review",
        "The record is restored, or its deletion eligibility and blockers are clearly displayed.",
        "Permanent deletion may be irreversible. Never proceed solely to remove an unwanted list entry; follow retention and clinic authorization rules.",
        True,
    ),
    proc(
        "Create and Manage Database Backups",
        "Create, schedule, verify, retain, and download system backups.",
        [
            "Select Database Backup from the sidebar.",
            "Review backup status, automatic schedule, file retention, verification status, and Backup History.",
            "To create a manual backup, select Create Backup Now and monitor progress until completion.",
            "In Automatic Backup Settings, enable or disable scheduling, select the frequency, set completed backups to retain, and save.",
            "For a completed backup, select Verify when available and review the verification result.",
            "Select Download only when an authorized secure copy is required, then protect the downloaded file.",
        ],
        "Database Backup status, settings, progress, and Backup History",
        "The backup operation or settings complete and the history shows file and verification status.",
        "Backup files may contain the entire database. Store, transfer, and dispose of them according to clinic security policy.",
        True,
    ),
    proc(
        "Run Integrity Checks and Safe Auto-Fix",
        "Detect data consistency issues and correct only eligible records.",
        [
            "Select Integrity Tools from the sidebar.",
            "Select Run Checks and wait for all categories to finish.",
            "Review Pass, No Issues, warning, and affected-record results.",
            "Expand a failed or warning check to inspect the affected records.",
            "For an issue explicitly marked Safe Auto-Fix, review the proposed change and confirm the fix.",
            "Rerun the affected check and verify the result.",
        ],
        "Integrity Tools check groups, affected records, and Safe Auto-Fix result",
        "Integrity results are refreshed and eligible corrected records pass the affected check.",
        "Do not manually alter database records from outside NgitiFy. Escalate issues that are not marked Safe Auto-Fix.",
        True,
    ),
]


EMR_READONLY = proc(
    "Review a Patient EMR in Read-Only Mode",
    "View authorized clinical information without changing the record.",
    [
        "Open Manage Patients and find the correct patient.",
        "Select View Patient EMR.",
        "Confirm the patient's identity.",
        "Review the Patient Profile, Medical and Dental History, Treatment History, Odontogram, and Radiograph Images tabs as permitted.",
        "Use Export PDF only for an authorized purpose and protect the exported file.",
    ],
    "Read-only EMR tabs and Export PDF action",
    "The authorized patient record is displayed without clinical editing controls.",
    "Clinical updates must be performed by an authorized dentist through the dentist-scoped workflow.",
)


SECRETARY_CHECKIN = proc(
    "Check In an Arriving Patient",
    "Move a confirmed patient into the in-clinic workflow.",
    [
        "On the Front Desk Dashboard or Schedule, locate the arriving patient's confirmed appointment.",
        "Verify the patient's name, appointment time, dentist, and procedure.",
        "Select Check In Patient or Mark as In Clinic.",
        "Confirm the action.",
        "Verify that the dashboard, schedule, and queue summary show the updated state.",
    ],
    "Today's appointments with Check In Patient and updated in-clinic status",
    "The appointment is marked In Clinic and becomes visible in the operational queue view.",
    "Confirm patient identity before check-in to avoid changing the wrong appointment.",
    True,
)


ROLE_CONTENT = {
    "System Administrator": {
        "scope": "System-wide administration across branches, users, appointments, patient management, inventory, configuration, audit, archive, backup, and integrity controls.",
        "access": [
            ("Scope", "All authorized branches and system administration records"),
            ("Core workspace", "Dashboard, Schedule, Manage Staffs, Manage Patients, Branches, Inventory"),
            ("Governance", "Activity Logs, Audit Trail, Roles and Permissions, Archive Review"),
            ("Platform controls", "System Config, Database Backup, Integrity Tools"),
            ("Clinical boundary", "May view authorized EMRs; clinical editing remains a dentist responsibility unless separately signed in with an authorized dentist workflow"),
        ],
        "sections": [
            ("Getting Started and Account Security", COMMON),
            ("Administrator Dashboard", [proc(
                "Review the Administrator Dashboard",
                "Monitor clinic-wide patients, staff, stock alerts, trends, activity, and appointments.",
                [
                    "Open Dashboard.",
                    "Review Active Patients, Total Staff, and Low Stock Alerts.",
                    "Review Patient Volume Trend, Treatment Breakdown, Recent System Activity, calendar, and appointment panels.",
                    "Select a card or related sidebar module when investigation or action is required.",
                ],
                "Administrator Dashboard statistic cards and overview panels",
                "The administrator can identify clinic-wide items requiring attention.",
            )]),
            ("Schedule Management", SCHEDULE),
            ("Staff and Patient Accounts", STAFF_MANAGEMENT + [PATIENT_SEARCH] + PATIENT_ADMIN + PATIENT_LIFECYCLE + [EMR_READONLY]),
            ("Branches and Analytics", BRANCHES),
            ("Inventory Management", [INVENTORY_VIEW] + INVENTORY_EDIT),
            ("System Governance and Configuration", ADMIN_ONLY),
        ],
    },
    "Owner-Dentist": {
        "scope": "Clinic-wide owner management plus dentist clinical functions for an owner account explicitly enabled as a dentist.",
        "access": [
            ("Owner scope", "Clinic-wide Dashboard, Schedule, staff, patients, branches, inventory, notifications, and activity"),
            ("Dentist scope", "My Schedule, My Patients, editable EMR, odontogram, radiograph review, and material usage"),
            ("Important distinction", "Use Manage Patients for organization-wide management; use My Patients for assigned clinical work"),
            ("Not included", "Administrator-only system configuration, audit trail, archive review, backup, and integrity tools"),
        ],
        "sections": [
            ("Getting Started and Account Security", COMMON),
            ("Owner Dashboard", [proc(
                "Review the Owner Dashboard",
                "Monitor clinic-wide appointments, patients, staff, activity, and supply health.",
                [
                    "Open Dashboard.",
                    "Review today's appointments, total patients, active staff, alerts, activity snapshot, calendar, and operational charts.",
                    "Select a statistic card to open Schedule, Manage Patients, or Manage Staffs.",
                    "Follow up in the relevant module before making a management decision.",
                ],
                "Owner Dashboard statistic cards and clinic-wide panels",
                "The owner can identify operational items requiring attention.",
            )]),
            ("Clinic-Wide and Personal Schedule", SCHEDULE + [proc(
                "Use My Schedule",
                "Focus on appointments assigned to the Owner-Dentist.",
                [
                    "Select My Schedule from the sidebar.",
                    "Review entries assigned to the signed-in Owner-Dentist.",
                    "Use search, filter, view, edit, status, and completion actions as appropriate.",
                    "Open the related patient EMR when clinical context is required.",
                ],
                "Owner-Dentist My Schedule view",
                "Only the Owner-Dentist's scoped schedule is shown for clinical work.",
            )]),
            ("Staff and Patient Management", STAFF_MANAGEMENT + [PATIENT_SEARCH] + PATIENT_ADMIN + PATIENT_LIFECYCLE + [EMR_READONLY]),
            ("My Patients and Clinical Records", CLINICAL),
            ("Branches and Analytics", BRANCHES),
            ("Inventory Management", [INVENTORY_VIEW] + INVENTORY_EDIT),
        ],
    },
    "Branch Manager": {
        "scope": "Operational management for the manager's assigned branch, subject to the current role permission matrix.",
        "access": [
            ("Scope", "Assigned branch staff, patients, schedule, analytics, inventory, notifications, and activity"),
            ("Staff", "Dentists and secretaries within the assigned branch"),
            ("Patient records", "Management actions according to permission; clinical EMR review is read-only"),
            ("Permissions", "A module may be full access, read-only, or unavailable based on administrator configuration"),
        ],
        "sections": [
            ("Getting Started and Account Security", COMMON),
            ("Branch Dashboard", [proc(
                "Review the Branch Manager Dashboard",
                "Monitor the assigned branch's appointments, patients, staff, activity, and queue operations.",
                [
                    "Open Dashboard and confirm the displayed branch label.",
                    "Review today's appointments, patient count, staff count, alerts, activity snapshot, calendar, and branch panels.",
                    "Select a statistic card to open the related branch-scoped module.",
                    "Investigate discrepancies before taking action.",
                ],
                "Branch Manager Dashboard with assigned branch label",
                "The manager sees the operational overview for the assigned branch.",
            )]),
            ("Branch Schedule", SCHEDULE),
            ("Branch Staff and Patients", STAFF_MANAGEMENT + [PATIENT_SEARCH] + PATIENT_ADMIN + PATIENT_LIFECYCLE + [EMR_READONLY]),
            ("Branch Analytics", [BRANCHES[2]]),
            ("Branch Inventory", [INVENTORY_VIEW] + INVENTORY_EDIT),
        ],
    },
    "Dentist": {
        "scope": "Assigned appointments, assigned patients, clinical EMR work, radiograph review, odontogram, treatment records, and material usage.",
        "access": [
            ("Scope", "Dentist-assigned schedule and patients according to branch assignment"),
            ("Clinical records", "Medical and dental history, treatment logs, odontogram, and radiographs"),
            ("Inventory", "Material selection through Material Usage; general inventory editing is not part of the dentist workflow"),
            ("Clinical responsibility", "AI suggestions and generated summaries require independent dentist review"),
        ],
        "sections": [
            ("Getting Started and Account Security", COMMON),
            ("Dentist Dashboard", [proc(
                "Review the Dentist Dashboard",
                "Monitor today's appointments, assigned patients, material logs, recent activity, calendar, and clinical follow-ups.",
                [
                    "Open Dashboard.",
                    "Review today's appointments, patient count, this month's material logs, recent account activity, calendar, and follow-up panels.",
                    "Open an appointment or patient record requiring attention.",
                    "Confirm source data in Schedule or the EMR before documenting care.",
                ],
                "Dentist Dashboard statistic cards and clinical panels",
                "The dentist can prioritize assigned clinical work.",
            )]),
            ("Dentist Schedule", SCHEDULE),
            ("Patients and Clinical Records", [PATIENT_SEARCH] + CLINICAL),
        ],
    },
    "Secretary": {
        "scope": "Front-desk operations for the assigned branch: appointments, check-in, patient registration and profile maintenance, read-only EMR access, notifications, and activity.",
        "access": [
            ("Scope", "Assigned branch schedule and patients"),
            ("Front desk", "Appointments, phone-call bookings, walk-ins, check-in, and status coordination"),
            ("Patient records", "Add and edit patient profile details; clinical EMR is read-only"),
            ("Not included", "Clinical documentation, staff management, branch configuration, inventory editing, or system administration"),
        ],
        "sections": [
            ("Getting Started and Account Security", COMMON),
            ("Front Desk Dashboard", [proc(
                "Review the Front Desk Dashboard",
                "Monitor today's appointments, patient count, queue, activity, calendar, and arrivals.",
                [
                    "Open Dashboard.",
                    "Review today's appointments, total patients, queue summary, alerts, activity snapshot, and calendar.",
                    "Open the relevant appointment or patient module from a statistic card.",
                    "Use the arrival action only after verifying the patient.",
                ],
                "Front Desk Dashboard statistic cards and today's appointments",
                "The secretary can prioritize arrivals and front-desk tasks.",
            ), SECRETARY_CHECKIN]),
            ("Appointment and Queue Operations", SCHEDULE),
            ("Patient Registration and Records", [PATIENT_SEARCH] + PATIENT_ADMIN + [PATIENT_LIFECYCLE[1], EMR_READONLY]),
        ],
    },
}


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            cell.width = Inches(widths[min(idx, len(widths) - 1)] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run(run, size=None, color=DARK, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_box(paragraph, fill, border):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "8")
        node.set(qn("w:color"), border)
        p_bdr.append(node)
    p_pr.append(p_bdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for item in (fld_begin, instr, fld_sep, fld_text, fld_end):
        run._r.append(item)
    set_run(run, size=9, color=MID_GRAY)


def new_step_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = getattr(doc, "_ngitify_step_abstract_id", None)
    if abstract_id is None:
        existing = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(existing, default=-1) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "%1.")
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, ind])
        level.extend([start, num_fmt, lvl_text, suff, p_pr])
        abstract.append(level)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            first_num.addprevious(abstract)
        doc._ngitify_step_abstract_id = abstract_id

    existing_num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_step_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_document(doc, role):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 8),
        ("Subtitle", 14, MID_GRAY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = name.startswith("Heading") or name == "Title"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run(f"NGITIFY  |  {role.upper()} USER MANUAL")
    set_run(hr, size=8.5, color=MID_GRAY, bold=True)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Dentime Dental Clinic  |  Internal Use  |  Page ")
    set_run(fr, size=9, color=MID_GRAY)
    add_page_field(fp)


def add_cover(doc, role, scope):
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NGITIFY")
    set_run(r, size=13, color=CYAN, bold=True)
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{role}\nUser Manual")
    set_run(r, size=30, color=NAVY, bold=True)
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(scope)
    set_run(r, size=13, color=MID_GRAY)
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.right_indent = Inches(0.45)
    p.paragraph_format.space_after = Pt(40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEB APPLICATION  |  DESKTOP / LAPTOP")
    set_run(r, size=10, color=WHITE, bold=True)
    paragraph_box(p, BLUE, BLUE)
    p.paragraph_format.space_after = Pt(46)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Version 1.0  |  September 2026")
    set_run(r, size=10.5, color=MID_GRAY, bold=True)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared for Dentime Dental Clinic")
    set_run(r, size=10, color=MID_GRAY, italic=True)
    doc.add_page_break()


def add_note(doc, label, text, warning=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    r = p.add_run(f"{label}: ")
    set_run(r, size=10.5, color=RED if warning else BLUE, bold=True)
    r = p.add_run(text)
    set_run(r, size=10.5, color=DARK)
    paragraph_box(p, PALE_RED if warning else PALE_BLUE, RED if warning else CYAN)


def add_scope_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "AREA"
    headers[1].text = "ROLE ACCESS / RESPONSIBILITY"
    for cell in headers:
        shade_cell(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            set_run(run, size=9.5, color=WHITE, bold=True)
    for label, detail in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = detail
        shade_cell(cells[0], PALE_BLUE)
        for run in cells[0].paragraphs[0].runs:
            set_run(run, size=10, color=NAVY, bold=True)
        for run in cells[1].paragraphs[0].runs:
            set_run(run, size=10, color=DARK)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [2400, 6960])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot_placeholder(doc, description):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SCREENSHOT TO ADD\n")
    set_run(r, size=10, color=BLUE, bold=True)
    r = p.add_run(f"Capture: {description}\n\n\n")
    set_run(r, size=9.5, color=MID_GRAY, italic=True)
    paragraph_box(p, PALE_GRAY, "A9BAC6")


def add_procedure(doc, number, item):
    p = doc.add_paragraph(style="Heading 2")
    p.add_run(f"{number} {item['title']}")

    p = doc.add_paragraph()
    r = p.add_run("Purpose: ")
    set_run(r, size=11, color=NAVY, bold=True)
    r = p.add_run(item["goal"])
    set_run(r, size=11, color=DARK)

    p = doc.add_paragraph()
    r = p.add_run("Steps")
    set_run(r, size=10.5, color=BLUE, bold=True)
    p.paragraph_format.space_after = Pt(3)
    num_id = new_step_numbering_id(doc)
    for step in item["steps"]:
        p = doc.add_paragraph(step, style="List Number")
        apply_step_numbering(p, num_id)
        p.paragraph_format.keep_together = True

    p = doc.add_paragraph()
    r = p.add_run("Expected result: ")
    set_run(r, size=10.5, color=GREEN, bold=True)
    r = p.add_run(item["result"])
    set_run(r, size=10.5, color=DARK)
    p.paragraph_format.keep_together = True

    if item.get("note"):
        add_note(doc, "Warning" if item.get("warning") else "Note", item["note"], item.get("warning", False))
    add_screenshot_placeholder(doc, item["screenshot"])


def add_front_matter(doc, role, content):
    doc.add_heading("About This Manual", level=1)
    doc.add_paragraph(
        f"This manual explains how the {role} uses the NgitiFy web application. It is based on the current role routing, interface controls, and tested workflows in the project as of September 2026. The instructions are intended for desktop or laptop use."
    )
    add_note(
        doc,
        "Permission note",
        "Visible menus and available actions may differ when the System Administrator changes the role permission matrix, when the account is limited to a branch, or when a feature toggle is disabled.",
    )
    add_note(
        doc,
        "Screenshot workflow",
        "Each procedure includes a labeled screenshot area. Replace that area with a current screenshot from the same role account and keep confidential patient or staff information hidden in training copies.",
    )

    doc.add_heading("Role Scope at a Glance", level=1)
    add_scope_table(doc, content["access"])

    doc.add_heading("Requirements and Safe Use", level=1)
    requirements = [
        "A desktop or laptop with a current version of Chrome, Edge, Firefox, or Safari.",
        "A stable network connection and the clinic's official NgitiFy web address.",
        f"An active {role} account with the correct branch assignment and permissions.",
        "Access to the registered email account for activation, verification, or password recovery.",
        "A private workspace when viewing patient, staff, financial, audit, or backup information.",
    ]
    for req in requirements:
        doc.add_paragraph(req, style="List Bullet")
    add_note(doc, "Security", "Do not share credentials, verification codes, exported reports, patient files, or backup files. Confirm the person and record before any status, clinical, financial, inventory, archive, or permission change.", True)

    doc.add_heading("Contents at a Glance", level=1)
    for index, (section_title, procedures) in enumerate(content["sections"], 1):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{index}. {section_title} ({len(procedures)} procedures)")
        set_run(r, size=10.5, color=DARK)
    doc.add_page_break()


def add_appendix(doc, role):
    doc.add_page_break()
    doc.add_heading("Quick Troubleshooting", level=1)
    items = [
        ("Cannot sign in", "Confirm the registered email, password, account status, and internet connection. Use Forgot Password if necessary."),
        ("Activation or OTP email did not arrive", "Check spam or junk folders, confirm the email address, and request only one new message before trying again."),
        ("A menu or button is missing", "The feature may be disabled, branch-scoped, read-only, or unavailable under the current permission matrix. Contact the System Administrator."),
        ("A record is not listed", "Clear filters, confirm the branch and status, check spelling, and verify that the user is authorized for the record."),
        ("A save action fails", "Review required fields, validation messages, duplicate warnings, file limits, and connection status before resubmitting."),
        ("The session expires", "Sign in again. NgitiFy may log users out after the configured inactivity period."),
    ]
    for title, detail in items:
        p = doc.add_paragraph()
        r = p.add_run(f"{title}: ")
        set_run(r, size=10.5, color=NAVY, bold=True)
        r = p.add_run(detail)
        set_run(r, size=10.5, color=DARK)

    doc.add_heading("Before Publishing This Manual", level=1)
    checks = [
        "Insert current screenshots captured with a non-production or safely redacted account.",
        "Verify the clinic web address, branch names, procedures, time slots, and enabled features.",
        "Confirm role permissions against the production permission matrix.",
        "Check all destructive-action and privacy warnings with clinic policy.",
        "Update the version and revision date after any interface or workflow change.",
    ]
    for check in checks:
        doc.add_paragraph(check, style="List Bullet")
    add_note(doc, "Document control", f"Role: {role} | Manual version: 1.0 | Revision: September 2026 | Product: NgitiFy", False)


def build_manual(role, content):
    doc = Document()
    configure_document(doc, role)
    add_cover(doc, role, content["scope"])
    add_front_matter(doc, role, content)

    for section_index, (section_title, procedures) in enumerate(content["sections"], 1):
        doc.add_heading(f"{section_index}. {section_title}", level=1)
        for procedure_number, item in enumerate(procedures, 1):
            add_procedure(doc, f"{section_index}.{procedure_number}", item)
        if section_index < len(content["sections"]):
            doc.add_page_break()

    add_appendix(doc, role)
    out = OUT_DIR / ROLE_FILES[role]
    doc.core_properties.title = f"NgitiFy {role} User Manual"
    doc.core_properties.subject = f"Role-based web application manual for {role}"
    doc.core_properties.author = "Dentime Dental Clinic"
    doc.core_properties.keywords = "NgitiFy, user manual, role-based access, Dentime"
    doc.core_properties.comments = "Generated from the current NgitiFy role routing and operational workflows."
    doc.save(out)
    return out


if __name__ == "__main__":
    for role, content in ROLE_CONTENT.items():
        print(build_manual(role, content).resolve())
