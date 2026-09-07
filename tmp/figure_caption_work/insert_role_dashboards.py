from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\Administrator\Documents\4.4 DESCRIPTION OF THE SYSTEM - Captioned with Public Website.docx")
OUTPUT = Path(r"C:\Users\Administrator\Documents\4.4 DESCRIPTION OF THE SYSTEM - Captioned with Role Dashboards.docx")
IMAGE_DIR = Path(r"C:\Users\Administrator\Desktop\NGITIFY DENTIME\tmp\figure_caption_work\role_dashboards")

titles = [
    "Owner Dashboard 1 (Web)",
    "Owner Dashboard 2 (Web)",
    "Branch Manager Dashboard 1 (Web)",
    "Branch Manager Dashboard 2 (Web)",
    "Dentist Dashboard 1 (Web)",
    "Dentist Dashboard 2 (Web)",
    "Front Desk Dashboard 1 (Web)",
    "Front Desk Dashboard 2 (Web)",
    "Patient Dashboard 1 (Web)",
    "Patient Dashboard 2 (Web)",
    "Recommended Visit Window Explanation (Web)",
]
images = [IMAGE_DIR / f"role-dashboard-{i:02d}.png" for i in range(1, 12)]
for image in images:
    if not image.exists():
        raise FileNotFoundError(image)

descriptions = {
    16: (
        "Figures 15 and 16 show the owner dashboard, which provides a clinic-wide overview of appointments, "
        "patient records, active staff, tracked branches, unread notifications, low-stock items, recent activity, "
        "and the latest operational alerts."
    ),
    18: (
        "Figures 17 and 18 show the branch manager dashboard. It summarizes appointments, patients, branch staff, "
        "the live queue, notifications, stock alerts, branch activity, the clinic calendar, and recent branch alerts."
    ),
    20: (
        "Figures 19 and 20 show the dentist dashboard, where a dentist can review the daily schedule, assigned "
        "patients, material usage, unread alerts, recent account activity, the monthly calendar, and today's appointments."
    ),
    22: (
        "Figures 21 and 22 show the front desk dashboard. It presents the branch schedule, registered patients, live "
        "queue, unread notifications, pending confirmations, new registrations, and shortcuts for patient and appointment management."
    ),
    25: (
        "Figures 23 through 25 show the patient dashboard and its visit-guidance explanation. Patients can review the "
        "calendar, selected-day schedule, appointments, daily Oral Health Management status, and recommended visit window, "
        "including the clinic, system, and recent oral-health information used to determine the guidance."
    ),
}


def set_font(run, name="Times New Roman", size=Pt(12), bold=False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def replace_text(paragraph, text, *, align=None, indent=False):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    set_font(paragraph.add_run(text))
    if align is not None:
        paragraph.alignment = align
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0


def shift_references(text):
    def repl(match):
        value = int(match.group())
        return str(value + 11) if value >= 15 else str(value)
    return re.sub(r"(?<![A-Za-z])\d+(?![A-Za-z])", repl, text)


doc = Document(SOURCE)

# Rename the two existing administrator sidebar figures as requested.
for p in doc.paragraphs:
    if p.text.strip().startswith("Figure 13."):
        replace_text(p, "Figure 13. Administrator Dashboard 3 – Sidebar Expanded (Web)", align=WD_ALIGN_PARAGRAPH.CENTER)
    elif p.text.strip().startswith("Figure 14."):
        replace_text(p, "Figure 14. Administrator Dashboard 4 – Sidebar Collapsed (Web)", align=WD_ALIGN_PARAGRAPH.CENTER)

# Renumber all existing captions from Figure 15 onward to make room for the
# eleven role-dashboard figures.
caption_elements = set()
for p in doc.paragraphs:
    match = re.fullmatch(r"Figure (\d+)\. (.+)", p.text.strip())
    if not match:
        continue
    number = int(match.group(1))
    title = match.group(2)
    if number >= 15:
        replace_text(p, f"Figure {number + 11}. {title}", align=WD_ALIGN_PARAGRAPH.CENTER)
    caption_elements.add(p._p)

# Update the administrator dashboard/sidebar explanation and shift every later
# narrative figure reference by the same offset.
admin_description = None
for p in doc.paragraphs:
    if p._p in caption_elements:
        continue
    text = p.text.strip()
    if text.startswith("Figures 11, 12, 13, and 14"):
        admin_description = p
        replace_text(
            p,
            "Figures 11 and 12 show the administrator dashboard and its operational summaries, charts, calendar, recent activity, and appointments. Figures 13 and 14 show the system sidebar in its expanded and collapsed states, providing full-label or compact-icon navigation according to the user's preference.",
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            indent=True,
        )
    elif re.match(r"^Figures? \d", text):
        replace_text(p, shift_references(text), align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)

if admin_description is None:
    raise RuntimeError("Administrator dashboard description was not found")

# Insert the role dashboards before the existing NgitifyBot figure.
from docx.text.paragraph import Paragraph
node = admin_description._p.getnext()
anchor = None
while node is not None:
    if node.xpath('.//w:drawing | .//w:pict'):
        anchor = Paragraph(node, admin_description._parent)
        break
    node = node.getnext()
if anchor is None:
    raise RuntimeError("Could not locate the figure following the administrator dashboard sequence")

for offset, (image_path, title) in enumerate(zip(images, titles)):
    number = 15 + offset
    image_p = anchor.insert_paragraph_before()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.keep_with_next = True
    if number in {15, 17, 19, 21, 23, 25}:
        image_p.paragraph_format.page_break_before = True
    width = 6.25
    image_p.add_run().add_picture(str(image_path), width=Inches(width))

    caption_p = anchor.insert_paragraph_before()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(3)
    caption_p.paragraph_format.space_after = Pt(3)
    set_font(caption_p.add_run(f"Figure {number}. {title}"))

    if number in descriptions:
        description_p = anchor.insert_paragraph_before()
        description_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        description_p.paragraph_format.first_line_indent = Inches(0.5)
        description_p.paragraph_format.space_before = Pt(6)
        description_p.paragraph_format.space_after = Pt(6)
        description_p.paragraph_format.line_spacing = 1.0
        set_font(description_p.add_run(descriptions[number]))

# Resume the remainder of the system description on a fresh page.
anchor.paragraph_format.page_break_before = True

while doc.paragraphs and not doc.paragraphs[-1].text.strip() and not doc.paragraphs[-1]._p.xpath('.//w:drawing | .//w:pict'):
    trailing = doc.paragraphs[-1]._p
    trailing.getparent().remove(trailing)

doc.save(OUTPUT)
print(f"Saved {OUTPUT}")
print("Inserted 11 role-dashboard figures and renumbered the complete sequence to Figure 198")
